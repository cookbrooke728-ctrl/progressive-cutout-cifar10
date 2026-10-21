"""Run with the bundled document runtime after the experiment audit succeeds."""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from content import make_pages

ROOT = Path(__file__).resolve().parents[1]


def set_font(run, size=10.5, bold=False, east="宋体"):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), east)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")


def configure_styles(doc):
    # Remove decorative borders inherited from the runtime's default template.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for border in doc.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    section.footer_distance = Cm(.8)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Pt(21)
    for name, size in (("Heading 1", 13), ("Heading 2", 11.5)):
        style = doc.styles[name]
        style.font.size, style.font.bold = Pt(size), True
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_text(doc, text, size=10.5, style=None, center=False, indent=True, spacing=1.4):
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(21 if indent and not center else 0)
    paragraph.paragraph_format.line_spacing = spacing
    set_font(paragraph.add_run(text), size)
    return paragraph


def add_table(doc, caption, headers, rows, widths):
    cap = add_text(doc, caption, 9, center=True, indent=False, spacing=1.1)
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(5)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths is None:
        widths = [15.7 / len(headers)] * len(headers)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    for r, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if r == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                element = OxmlElement(f"w:{edge}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    element.set(qn(f"w:{key}"), value)
                borders.append(element)
            tc.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:w"), "65")
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            tc.append(margins)
            if r == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EEF1F4")
                tc.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(headers) > 2 or i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_font(run, 9 if len(rows) > 8 else 9.5, r == 0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = .2
    set_font(spacer.add_run(""), 2)


def main():
    summary_path = ROOT / "output/summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    audit = json.loads((ROOT / "output/audit.json").read_text(encoding="utf-8"))
    if audit["complete_runs"] != 12:
        raise RuntimeError("Paper generation requires 12 audited formal runs")
    publication = json.loads((ROOT / "paper/publication.json").read_text(encoding="utf-8"))
    pages = make_pages(summary, publication)
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "渐进式随机遮挡对图像分类性能的影响"
    doc.core_properties.subject = "CIFAR-10 Cutout reproduction and controlled progressive probability study"
    doc.core_properties.author = publication.get("author", "")
    markdown = []
    for index, blocks in enumerate(pages):
        if index:
            doc.add_page_break()
        for block in blocks:
            kind = block[0]
            if kind in ("title", "subtitle", "en_title", "en_subtitle", "center"):
                size = {"title": 17, "subtitle": 12, "en_title": 13, "en_subtitle": 11, "center": 10}[kind]
                paragraph = add_text(doc, block[1], size, style="Title" if kind == "title" else None, center=True, indent=False, spacing=1.1)
                paragraph.paragraph_format.space_after = Pt(4 if kind == "center" else 7)
                if kind in ("title", "en_title"):
                    paragraph.runs[0].bold = True
                markdown.extend([("# " if kind == "title" else "") + block[1], ""])
            elif kind in ("abstract", "abstract_en"):
                paragraph = add_text(doc, "", 9.5, indent=False, spacing=1.1)
                set_font(paragraph.add_run(block[1] + ("：" if kind == "abstract" else ": ")), 9.5, True)
                set_font(paragraph.add_run(block[2]), 9.5)
                markdown.extend([f"**{block[1]}** {block[2]}", ""])
            elif kind in ("h1", "h2"):
                paragraph = doc.add_paragraph(block[1], style="Heading 1" if kind == "h1" else "Heading 2")
                for run in paragraph.runs:
                    set_font(run, 13 if kind == "h1" else 11.5, True)
                markdown.extend([("## " if kind == "h1" else "### ") + block[1], ""])
            elif kind in ("p", "reference", "small"):
                paragraph = add_text(doc, block[1], 9 if kind in ("reference", "small") else 10.5,
                                     indent=kind == "p", spacing=1.2 if kind != "p" else 1.4)
                if kind == "reference":
                    paragraph.paragraph_format.space_after = Pt(6)
                markdown.extend([block[1], ""])
            elif kind == "eq":
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Pt(0)
                math_para, math = OxmlElement("m:oMathPara"), OxmlElement("m:oMath")
                run, text = OxmlElement("m:r"), OxmlElement("m:t")
                text.text = block[1]
                run.append(text)
                math.append(run)
                math_para.append(math)
                paragraph._p.append(math_para)
                markdown.extend([block[1], ""])
            elif kind == "figure":
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run()
                shape = run.add_picture(str(ROOT / "output/figures" / block[1]), width=Cm(block[3]))
                shape._inline.docPr.set("descr", block[2])
                cap = add_text(doc, block[2], 9, center=True, indent=False, spacing=1.1)
                cap.paragraph_format.space_after = Pt(7)
                markdown.extend([f"![{block[2]}](figures/{block[1]})", ""])
            elif kind == "table":
                add_table(doc, *block[1:])
                markdown.extend([block[1], "", "|" + "|".join(block[2]) + "|", "|" + "|".join(["---"] * len(block[2])) + "|"])
                markdown.extend("|" + "|".join(map(str, row)) + "|" for row in block[3])
                markdown.append("")
    destination = ROOT / "output/paper"
    destination.mkdir(parents=True, exist_ok=True)
    path = destination / "progressive_cutout_paper.docx"
    doc.save(path)
    (ROOT / "output/paper.md").write_text("\n".join(markdown), encoding="utf-8")
    import hashlib
    manifest = {"summary_sha256": hashlib.sha256(summary_path.read_bytes()).hexdigest(),
                "planned_pages": len(pages), "author_pending": "待填写" in publication["author"],
                "github_pending": not bool(publication["repository_url"] and publication["release_url"]),
                "word_count_characters": len("\n".join(markdown))}
    (destination / "inputs.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(json.dumps({"docx": str(path), **manifest}))


if __name__ == "__main__":
    main()
